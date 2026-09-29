from docx import Document
from pathlib import Path

md_path = Path(r"d:\PoomWork\RAE67_bio_ins-main\lab_02\results\student_comparison\lab2_report.md")
out_path = Path(r"d:\PoomWork\RAE67_bio_ins-main\lab_02\results\student_comparison\lab2_report.docx")

doc = Document()

with md_path.open("r", encoding="utf-8") as f:
    lines = f.read().splitlines()

for line in lines:
    if not line.strip():
        continue
    if line.startswith("# "):
        doc.add_heading(line[2:], level=1)
    elif line.startswith("## "):
        doc.add_heading(line[3:], level=2)
    elif line.startswith("### "):
        doc.add_heading(line[4:], level=3)
    elif line.startswith("- "):
        doc.add_paragraph(line[2:])
    elif line.startswith("|"):
        doc.add_paragraph(line)
    else:
        doc.add_paragraph(line)

doc.save(out_path)
print(f"Saved Word file: {out_path}")
